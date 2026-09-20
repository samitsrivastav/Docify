"""PPTX text + structure extraction using python-pptx."""
import io
import logging
from typing import List, Union

from pptx import Presentation

from app.parser.pdf_parser import DocumentParseError, PageContent, ParsedDocument

logger = logging.getLogger(__name__)


def parse_pptx(file_source: Union[bytes, str], filename: str) -> ParsedDocument:
    """Parse a PPTX from raw bytes or a file path. One 'page' == one slide."""
    try:
        source = io.BytesIO(file_source) if isinstance(file_source, (bytes, bytearray)) else file_source
        prs = Presentation(source)
    except Exception as e:
        raise DocumentParseError(f"Could not open presentation '{filename}': {e}") from e

    if len(prs.slides) == 0:
        raise DocumentParseError(f"'{filename}' has no slides.")

    pages: List[PageContent] = []
    total_words = 0

    for i, slide in enumerate(prs.slides):
        texts: List[str] = []
        headings: List[str] = []

        title_shape = slide.shapes.title
        title_text = ""
        if title_shape is not None and title_shape.has_text_frame:
            title_text = title_shape.text.strip()
        if title_text:
            headings.append(title_text)
            texts.append(title_text)

        for shape in slide.shapes:
            if shape is title_shape or not shape.has_text_frame:
                continue
            shape_text = "\n".join(
                p.text for p in shape.text_frame.paragraphs if p.text.strip()
            ).strip()
            if shape_text:
                texts.append(shape_text)

        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                texts.append(f"Notes: {notes_text}")

        full_text = "\n".join(texts).strip()
        pages.append(PageContent(page_number=i + 1, text=full_text, headings=headings))
        total_words += len(full_text.split())

    if total_words == 0:
        raise DocumentParseError(f"'{filename}' contains no extractable text.")

    return ParsedDocument(
        filename=filename,
        doc_type="pptx",
        pages=pages,
        page_count=len(pages),
        word_count=total_words,
    )
