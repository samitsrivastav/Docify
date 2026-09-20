import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.parser import DocumentParseError
from app.parser.pptx_parser import parse_pptx


def make_pptx_bytes(slide_titles):
    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    for title in slide_titles:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.text = f"Body content for {title}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_parse_pptx_extracts_slides():
    pptx_bytes = make_pptx_bytes(["Intro", "Details"])
    parsed = parse_pptx(pptx_bytes, "deck.pptx")
    assert parsed.doc_type == "pptx"
    assert parsed.page_count == 2
    assert parsed.pages[0].headings == ["Intro"]
    assert "Body content for Intro" in parsed.pages[0].text


def test_parse_pptx_raises_on_no_slides():
    prs = Presentation()
    buf = io.BytesIO()
    prs.save(buf)
    with pytest.raises(DocumentParseError):
        parse_pptx(buf.getvalue(), "empty.pptx")
